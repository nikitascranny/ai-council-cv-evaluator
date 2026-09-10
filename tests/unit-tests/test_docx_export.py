"""The rewritten CV as a .docx.

The guard's verdict has to survive into the file itself - a document outlives
the page that offered it, and a rejected rewrite must never read as a sound one.
"""

import io
import zipfile
from types import SimpleNamespace

from docx import Document

from ui import docx_export


def _result(**kw):
    base = dict(candidate="Jane Doe", job_title="Secretary", agent_model="phi4",
                original_text="orig", rewritten_text="First block.\n\nSecond block.",
                used_fallback=False, rejected=False)
    base.update(kw)
    return SimpleNamespace(**base)


def _paragraphs(result):
    return [p.text for p in docx_export.build_document(result).paragraphs]


def test_the_bytes_are_a_real_docx():
    """A .docx is a zip. Asserting that catches a truncated or empty write
    that a paragraph-level check would sail past."""
    raw = docx_export.docx_bytes(_result())
    assert zipfile.is_zipfile(io.BytesIO(raw))
    assert Document(io.BytesIO(raw)).paragraphs


def test_the_rewritten_text_is_split_into_paragraphs_not_one_slab():
    assert "First block." in _paragraphs(_result())
    assert "Second block." in _paragraphs(_result())


def test_a_fallback_says_in_the_document_that_it_is_the_original_cv():
    """used_fallback means the guard rejected the rewrite twice and v2
    substituted the candidate's own text. Someone opening this file later
    must not read it as a rewrite that happened not to change much."""
    text = " ".join(_paragraphs(_result(used_fallback=True)))
    assert "not a rewrite" in text
    assert "original CV" in text


def test_a_clean_rewrite_carries_no_warning():
    text = " ".join(_paragraphs(_result()))
    assert "not a rewrite" not in text
    assert "flagged" not in text


def test_a_flagged_but_kept_rewrite_says_so():
    text = " ".join(_paragraphs(_result(rejected=True)))
    assert "flagged" in text


def test_the_job_title_is_named_so_the_file_says_what_it_was_aimed_at():
    assert any("Secretary" in p for p in _paragraphs(_result()))


def test_the_filename_names_candidate_job_and_agent():
    """Two versions of one CV land in the same downloads folder, so the
    agent has to be in the name or the second silently overwrites nothing
    and the user cannot tell them apart."""
    name = docx_export.filename_for(_result())
    assert name.endswith(".docx")
    assert "Jane Doe" in name and "Secretary" in name and "phi4" in name


def test_a_colon_in_the_model_tag_does_not_reach_the_filename():
    """qwen2.5:14b is a valid Ollama tag. On NTFS a ":" in a filename opens
    an alternate data stream, and in a Content-Disposition header it breaks
    the parse."""
    name = docx_export.filename_for(_result(agent_model="qwen2.5:14b"))
    assert ":" not in name
    assert "qwen2.5" in name


def test_empty_rewritten_text_still_produces_a_readable_document():
    raw = docx_export.docx_bytes(_result(rewritten_text=""))
    assert zipfile.is_zipfile(io.BytesIO(raw))
