"""The rewritten CV as a Word document the candidate can edit.

rewriter_v2 returns plain text plus guard metadata and deliberately stops
there - its docstring says the export step exists so that data need not be
reshaped to build it. This is that step.

Pure by design: bytes in memory, no filesystem. The web UI serves the bytes
as a download; rewriter_v2's own CLI can write them to a file next to its
.txt and metadata.csv. Neither one is baked in here, and a test can assert on
the document without Ollama or a temp directory.

The fallback banner is written INTO the document rather than shown only on the
page that offered it. A .docx outlives the browser session that produced it,
and the whole point of the guard is that a rewrite it rejected must never be
mistaken for one it passed - least of all by whoever opens this file a week
later with no memory of the warning.
"""

import io

from docx import Document
from docx.shared import Pt

# The guard rejected this agent's rewrite twice, so rewritten_text is the
# candidate's ORIGINAL CV. Saying that plainly is the entire job of this
# paragraph - "fallback" is our word, not something a candidate would parse.
_FALLBACK_WARNING = (
    "This is your original CV, not a rewrite. The automated check found that "
    "the rewritten version added information your CV does not contain, so it "
    "was discarded rather than shown to you as if it were sound.")

_FLAGGED_WARNING = (
    "The automated check flagged this rewrite. Read it against your original "
    "before sending it anywhere.")


def _add_body(document, text: str):
    """One paragraph per blank-line-separated block of the rewritten text.

    Deliberately flat. rewritten_text is plain text with no structure a
    model reliably emits, so inventing headings here would be guesswork
    dressed as formatting. parser/cv_sections.py already does real heading
    detection and is the right basis for that later - it is another team
    member's module and a separate change.
    """
    for block in text.replace("\r\n", "\n").split("\n\n"):
        block = block.strip()
        if block:
            document.add_paragraph(block)


def build_document(result) -> "Document":
    """A Document for one agent's RewriteResult.

    Returned rather than saved so callers can serialise it however they need
    and tests can read the paragraphs back without touching a disk.
    """
    document = Document()
    document.styles["Normal"].font.size = Pt(11)

    heading = result.candidate or "Rewritten CV"
    document.add_heading(heading, level=1)
    if result.job_title:
        document.add_paragraph("Tailored for: {}".format(result.job_title))

    if result.used_fallback:
        document.add_paragraph(_FALLBACK_WARNING)
    elif result.rejected:
        document.add_paragraph(_FLAGGED_WARNING)

    _add_body(document, result.rewritten_text or "")
    return document


def docx_bytes(result) -> bytes:
    """The document as the bytes of a .docx file."""
    buffer = io.BytesIO()
    build_document(result).save(buffer)
    return buffer.getvalue()


def filename_for(result) -> str:
    """A download name that says whose CV, for what job, from which agent.

    Three CVs in a downloads folder all called "cv.docx" are useless, and the
    agent belongs in the name because the two versions differ - that is the
    whole reason the user was shown both.
    """
    parts = [result.candidate or "cv", result.job_title or "", result.agent_model or ""]
    keep = "-_(). "
    cleaned = []
    for part in parts:
        part = "".join(c if (c.isalnum() or c in keep) else "_" for c in part).strip()
        if part:
            cleaned.append(part)
    return " - ".join(cleaned).strip().strip(".") + ".docx"
